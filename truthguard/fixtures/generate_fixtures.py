from pathlib import Path

FIXTURES_DIR = Path(__file__).resolve().parent

def create_leave_policy_v1_txt():
    content = """ACME CORPORATION
POLICY MANUAL: LEAVE POLICY (VERSION 1.0)
Published: January 15, 2024
Status: Approved Policy
Category: Human Resources Regulations

Section 1: Purpose and Scope
This policy manual governs all permanent and contract personnel across all company branches.

Section 2: Casual Leave Entitlement
Employees are entitled to 12 casual leave days per year. Casual leaves must be scheduled with at least 24 hours advance notice to the department supervisor. Unused casual leaves expire at the end of each calendar year and cannot be carried forward.

Section 3: Sick Leave Entitlement
Employees receive 10 paid sick leave days annually. Any absence exceeding two consecutive business days requires verification from a licensed medical practitioner.

Section 4: Approval Process
Leave requests must be logged through the employee portal and approved by the reporting manager prior to taking time off.
"""
    file_path = FIXTURES_DIR / "leave_policy_v1.txt"
    with open(file_path, "w", encoding="utf-8") as f:
        f.write(content)
    print(f"Created: {file_path}")


def create_leave_amendment_v2_docx():
    import docx
    doc = docx.Document()
    doc.add_heading("ACME CORPORATION", level=0)
    doc.add_heading("POLICY AMENDMENT: CASUAL LEAVE REVISION (VERSION 2.0)", level=1)
    
    doc.add_paragraph("Published: March 10, 2025")
    doc.add_paragraph("Status: Executive Draft / Pending Board Ratification")
    doc.add_paragraph("Authority: Proposed Amendment")

    doc.add_heading("Section 1: Purpose of Amendment", level=2)
    doc.add_paragraph("This amendment proposes an enhancement to workplace flexibility and personal well-being entitlements for all staff.")

    doc.add_heading("Section 2: Revised Casual Leave Entitlement", level=2)
    doc.add_paragraph("Employees are entitled to 15 casual leave days per year.")

    doc.add_heading("Section 3: Effective Date and Governance Note", level=2)
    doc.add_paragraph(
        "This amendment is scheduled to supersede Leave Policy v1.0 contingent upon final Board of Directors sign-off on April 1, 2025. "
        "Until official corporate broadcast, previous active entitlements remain under review."
    )

    file_path = FIXTURES_DIR / "leave_amendment_v2.docx"
    doc.save(str(file_path))
    print(f"Created: {file_path}")


def create_employee_faq_pdf():
    import fitz  # PyMuPDF
    doc = fitz.open()
    page = doc.new_page(width=595, height=842) # A4 size

    text = """ACME CORPORATION — EMPLOYEE SELF-SERVICE FAQ
Published: June 2024
Category: Human Resources FAQ
Status: Informational FAQ

Frequently Asked Questions:

Q: How many casual leaves do employees receive?
A: Employees receive 12 casual leave days per year. Unused casual leaves cannot be encashed at year end.

Q: How do I submit an expense reimbursement?
A: Submit all itemized receipts within 30 days of the transaction via the finance portal.

Q: What is the probation period for new hires?
A: The standard probation period is 90 days from the official date of joining.
"""
    # Insert text into PDF page
    rect = fitz.Rect(50, 50, 545, 792)
    page.insert_textbox(rect, text, fontsize=11, fontname="helv", align=0)
    
    file_path = FIXTURES_DIR / "employee_faq.pdf"
    doc.save(str(file_path))
    doc.close()
    print(f"Created: {file_path}")


if __name__ == "__main__":
    create_leave_policy_v1_txt()
    create_leave_amendment_v2_docx()
    create_employee_faq_pdf()
    print("All 3 test fixtures successfully created!")
