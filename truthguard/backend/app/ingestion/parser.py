import io
from pathlib import Path
from typing import List, Tuple
from fastapi import HTTPException, status


class ExtractedSection:
    def __init__(self, heading: str, text: str, page_number: int = 1):
        self.heading = heading
        self.text = text
        self.page_number = page_number


def parse_pdf(file_bytes: bytes, filename: str) -> List[ExtractedSection]:
    """Extracts text and section headings from PDF using PyMuPDF (fitz)."""
    try:
        import fitz  # PyMuPDF
    except ImportError:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="PyMuPDF is not installed on the backend server."
        )

    try:
        doc = fitz.open(stream=file_bytes, filetype="pdf")
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Corrupt or unreadable PDF document '{filename}': {str(exc)}"
        )

    if doc.page_count == 0:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"PDF document '{filename}' is empty."
        )

    sections: List[ExtractedSection] = []
    total_text_len = 0

    for page_idx in range(doc.page_count):
        page = doc.load_page(page_idx)
        page_num = page_idx + 1
        raw_page_text = page.get_text("text")
        page_text = raw_page_text.strip() if isinstance(raw_page_text, str) else str(raw_page_text).strip()
        total_text_len += len(page_text)

        if not page_text:
            continue

        # Split page into logical sections by line headings
        lines = page_text.splitlines()
        current_heading = f"Page {page_num}"
        current_lines: List[str] = []

        for line in lines:
            trimmed = line.strip()
            if not trimmed:
                continue
            # Simple heading heuristic: starts with Section, Article, Q:, # or uppercase
            is_heading = (
                trimmed.lower().startswith(("section", "article", "chapter", "part")) or
                trimmed.startswith(("#", "Q:", "FAQ:")) or
                (len(trimmed) < 60 and trimmed.isupper() and len(trimmed) > 3)
            )
            if is_heading and current_lines:
                sec_text = "\n".join(current_lines).strip()
                if sec_text:
                    sections.append(ExtractedSection(current_heading, sec_text, page_num))
                current_heading = trimmed
                current_lines = [trimmed]
            else:
                if is_heading:
                    current_heading = trimmed
                current_lines.append(trimmed)

        if current_lines:
            sec_text = "\n".join(current_lines).strip()
            if sec_text:
                sections.append(ExtractedSection(current_heading, sec_text, page_num))

    doc.close()

    if total_text_len < 30:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"PDF document '{filename}' appears to be a scanned image or contains no extractable digital text."
        )

    return sections


def parse_docx(file_bytes: bytes, filename: str) -> List[ExtractedSection]:
    """Extracts text and headings from DOCX using python-docx."""
    try:
        import docx
    except ImportError:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="python-docx is not installed on the backend server."
        )

    try:
        doc = docx.Document(io.BytesIO(file_bytes))
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Corrupt or invalid DOCX document '{filename}': {str(exc)}"
        )

    sections: List[ExtractedSection] = []
    current_heading = "Document Header"
    current_lines: List[str] = []
    total_len = 0
    estimated_page = 1
    line_count = 0

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        total_len += len(text)
        line_count += 1
        if line_count % 35 == 0:
            estimated_page += 1

        style_name = para.style.name if (para.style and para.style.name) else ""
        is_heading = (
            "Heading" in style_name or
            "Title" in style_name or
            text.lower().startswith(("section", "article", "chapter", "part")) or
            (len(text) < 60 and text.isupper() and len(text) > 3)
        )

        if is_heading and current_lines:
            sec_text = "\n".join(current_lines).strip()
            if sec_text:
                sections.append(ExtractedSection(current_heading, sec_text, estimated_page))
            current_heading = text
            current_lines = [text]
        else:
            if is_heading:
                current_heading = text
            current_lines.append(text)

    # Also extract tables if present
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                current_lines.append(row_text)
                total_len += len(row_text)

    if current_lines:
        sec_text = "\n".join(current_lines).strip()
        if sec_text:
            sections.append(ExtractedSection(current_heading, sec_text, estimated_page))

    if total_len < 10:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"DOCX document '{filename}' is empty or has no readable text content."
        )

    return sections


def parse_txt(file_bytes: bytes, filename: str) -> List[ExtractedSection]:
    """Extracts text and headings from plain text files."""
    try:
        text = file_bytes.decode("utf-8")
    except UnicodeDecodeError:
        try:
            text = file_bytes.decode("latin-1")
        except Exception as exc:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Encoding error parsing text file '{filename}': {str(exc)}"
            )

    stripped = text.strip()
    if len(stripped) < 10:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Text file '{filename}' is empty or contains insufficient content."
        )

    sections: List[ExtractedSection] = []
    lines = stripped.splitlines()
    current_heading = "General"
    current_lines: List[str] = []
    estimated_page = 1
    line_count = 0

    for line in lines:
        trimmed = line.strip()
        if not trimmed:
            continue
        line_count += 1
        if line_count % 40 == 0:
            estimated_page += 1

        is_heading = (
            trimmed.lower().startswith(("section", "article", "chapter", "part")) or
            trimmed.startswith(("#", "==", "--", "Q:")) or
            (len(trimmed) < 60 and trimmed.isupper() and len(trimmed) > 3)
        )

        if is_heading and current_lines:
            sec_text = "\n".join(current_lines).strip()
            if sec_text:
                sections.append(ExtractedSection(current_heading, sec_text, estimated_page))
            current_heading = trimmed
            current_lines = [trimmed]
        else:
            if is_heading:
                current_heading = trimmed
            current_lines.append(trimmed)

    if current_lines:
        sec_text = "\n".join(current_lines).strip()
        if sec_text:
            sections.append(ExtractedSection(current_heading, sec_text, estimated_page))

    return sections


def parse_document(file_bytes: bytes, filename: str) -> Tuple[List[ExtractedSection], str]:
    """
    Dispatcher for PDF, DOCX, TXT.
    Returns (List[ExtractedSection], full_combined_text).
    """
    ext = Path(filename).suffix.lower().lstrip(".")
    if ext == "pdf":
        sections = parse_pdf(file_bytes, filename)
    elif ext == "docx":
        sections = parse_docx(file_bytes, filename)
    elif ext in ("txt", "text", "md"):
        sections = parse_txt(file_bytes, filename)
    else:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Unsupported file format '.{ext}'. Supported formats are: PDF, DOCX, TXT."
        )

    full_text = "\n\n".join(f"[{s.heading} - Page {s.page_number}]\n{s.text}" for s in sections)
    return sections, full_text
